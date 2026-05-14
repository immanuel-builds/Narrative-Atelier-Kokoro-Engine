from docx import Document
from pathlib import Path
from app.storage.markdown_handler import MarkdownHandler

class ExportService:
    @staticmethod
    def to_markdown(content: str, target_path: Path):
        target_path.write_text(content, encoding="utf-8")
        return str(target_path)

    @staticmethod
    def to_txt(content: str, target_path: Path):
        target_path.write_text(content, encoding="utf-8")
        return str(target_path)

    @staticmethod
    def to_docx(title: str, content: str, target_path: Path):
        doc = Document()
        doc.add_heading(title, 0)

        # Simple split by paragraphs
        for p in content.split('\n'):
            if p.strip():
                doc.add_paragraph(p.strip())

        doc.save(str(target_path))
        return str(target_path)

    @staticmethod
    def export_project(project_title: str, chapters_data: list, format: str = "md") -> str:
        # chapters_data: list of {"title": str, "content": str}
        safe_title = MarkdownHandler.sanitize_filename(project_title)
        export_dir = Path("Projects") / safe_title / "Exports"
        export_dir.mkdir(parents=True, exist_ok=True)

        full_content = f"# {project_title}\n\n"
        for i, ch in enumerate(chapters_data):
            full_content += f"## Chapter {i+1}: {ch['title']}\n\n{ch['content']}\n\n---\n\n"

        target_path = export_dir / f"{safe_title}_full.{format}"

        if format == "md":
            return ExportService.to_markdown(full_content, target_path)
        elif format == "txt":
            return ExportService.to_txt(full_content, target_path)
        elif format == "docx":
            return ExportService.to_docx(project_title, full_content, target_path)

        return ""
