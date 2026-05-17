from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pathlib import Path

class DOCXExporter:
    @staticmethod
    def export(title: str, author: str, chapters: list, target_path: Path):
        doc = Document()

        # Title Page
        title_para = doc.add_paragraph()
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title_para.add_run(title)
        run.font.size = Pt(24)
        run.bold = True

        author_para = doc.add_paragraph()
        author_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        author_para.add_run(f"by {author}")

        doc.add_page_break()

        # Chapters
        for i, ch in enumerate(chapters):
            doc.add_heading(f"Chapter {i+1}: {ch['title']}", level=1)
            for p in ch['content'].split('\n'):
                if p.strip():
                    doc.add_paragraph(p.strip())
            doc.add_page_break()

        doc.save(str(target_path))
        return str(target_path)
